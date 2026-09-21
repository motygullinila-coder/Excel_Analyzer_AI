import os

import pandas as pd
from openpyxl import load_workbook

class MetaAnalyzerFile:

    def __init__(self, path_file: str):
        self.path_file = path_file


    def get_workbook_info(self) -> dict:
        full_file_name = os.path.basename(self.path_file)
        extension = os.path.splitext(full_file_name)[1]

        workbook = load_workbook(filename=self.path_file, read_only=False)
        list_sheets = workbook.sheetnames

        return {
            'nameFile': full_file_name,
            'extension': extension,
            'sheetsCount': len(list_sheets)
        }


    def get_profile(self) -> dict:
        response_data = {}
        response_data['sheets'] = []

        workbook = load_workbook(self.path_file, read_only=False)

        list_sheets = workbook.sheetnames
        for element in list_sheets:
            title_list = element
            row_count = workbook[element].max_row
            col_count = workbook[element].max_column
            cells = row_count * col_count

            meta_list_data = {
                'nameList': title_list,
                'dimensions': {
                    'rows': row_count,
                    'columns': col_count,
                    'cells': cells
                }
            }

            response_data['sheets'].append(meta_list_data)

        return response_data


    def get_dataframe_info(self) -> dict:
        response_data = {}
        response_data['sheets'] = []

        file_excel = pd.ExcelFile(self.path_file)
        list_sheets = file_excel.sheet_names

        for element in list_sheets:
            dataframe = pd.read_excel(self.path_file, sheet_name=element)
            struct_response = {
                'nameList': element,
                'dimensions': {
                    'rows': dataframe.shape[0],
                    'columns': dataframe.shape[1],
                    'size': dataframe.size
                }
            }

            response_data['sheets'].append(struct_response)

        return response_data


    # -> Данный метод получает информацию о листе(получает 5 х 5 матрицу, где видно какие ячейки пустые, а какие имеют значение)
    def get_sheet_preview(self) -> dict:
        response_data = {}
        response_data['sheets'] = []

        workbook = load_workbook(filename=self.path_file, read_only=False)
        list_sheets = workbook.sheetnames

        for element_sheet in list_sheets:
            list_cell_value = []

            for row_element in workbook[element_sheet].iter_rows(min_row=1, max_row=5, min_col=1, max_col=5):
                row_list = []

                for cell_element in row_element:
                    row_list.append(cell_element.value)

                list_cell_value.append(row_list)

            struct_response = {
                'sheet': element_sheet,
                'preview': list_cell_value
            }

            response_data['sheets'].append(struct_response)
        return response_data
